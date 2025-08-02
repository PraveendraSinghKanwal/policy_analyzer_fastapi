from fastapi import APIRouter, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
from io import BytesIO
import zipfile
import openpyxl
from docx import Document
from reportlab.pdfgen import canvas
import os
import glob
from pathlib import Path
from .process_excel import process_excel_files

router = APIRouter()

def create_dummy_excel1() -> BytesIO:
    excel_buffer = BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ID", "Name", "Value","ID", "Name", "Value","ID", "Name", "Value","ID", "Name", "Value","ID", "Name", "Value","ID", "Name", "Value","ID", "Name", "Value","ID", "Name", "Value"])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    ws.append([1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123,1, "Item 1A", 123, 1, "Item 1A", 123,1, "Item 1A", 123])
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer

def create_dummy_excel2() -> BytesIO:
    excel_buffer = BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ID", "Name", "Value"])
    ws.append([1, "Item 2A", 234])
    ws.append([2, "Item 2B", 122])
    ws.append([3, "Item 2C", 655])
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer

def create_dummy_text() -> BytesIO:
    text_buffer = BytesIO()
    content = (
        "This is a dummy text file.\n"
        "Line 2: Example data.\n"
        "Line 3: More dummy content. the heart of the bustling city, \n"
    )
    text_buffer.write(content.encode("utf-8"))
    text_buffer.seek(0)
    return text_buffer

def create_dummy_doc() -> BytesIO:
    doc_buffer = BytesIO()
    document = Document()
    
    # Add paragraphs to the document
    paragraphs = [
        "This is a dummy DOC file.",
        "Line 2: Example data.",
        "Line 3: More dummy content."
    ]
    
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    
    # Save document to buffer
    document.save(doc_buffer)
    doc_buffer.seek(0)
    
    return doc_buffer

def create_dummy_pdf() -> BytesIO:
    pdf_buffer = BytesIO()
    pdf = canvas.Canvas(pdf_buffer)
    
    # Add lines of text
    lines = [
        "This is a dummy PDF file.",
        "Line 2: Example data.",
        "Line 3: More dummy content. The heart of the bustling city,"
    ]
    
    y = 800  # start from top of page (approx)
    for line in lines:
        pdf.drawString(100, y, line)
        y -= 20  # move down the page
    
    pdf.save()
    pdf_buffer.seek(0)
    
    return pdf_buffer


def add_file_to_zip(zip_file, filename: str, file_buffer: BytesIO):
    file_buffer.seek(0)
    zip_file.writestr(filename, file_buffer.read())



@router.post("/upload-pdf")
async def upload_pdf(file: UploadFile = File(...), local_folder: str = "excel_files"):
    # Accept any file, but you can keep PDF validation if needed
    try:
        import time
        # time.sleep(1)
        zip_buffer = BytesIO()
        with zipfile.ZipFile(zip_buffer, "w") as zip_file:
            # Add files from local folder
            add_folder_files_to_zip(zip_file, local_folder)
            
            # Continue with dummy files (existing functionality)
            # excel1 = create_dummy_excel1()
            # add_file_to_zip(zip_file, "Standard_Analyses_travell_policy.xlsx", excel1)
            # add_file_to_zip(zip_file, "Standard_Analyses_stay_policy.xlsx", excel1)
            # excel2 = create_dummy_excel2()
            # add_file_to_zip(zip_file, "Gap_Analyses_travell_policy.xlsx", excel2)
            # add_file_to_zip(zip_file, "Gap_Analyses_stay_policy.xlsx", excel2)
            # add_file_to_zip(zip_file, "Gap_Analyses_rest.xlsx", excel2)
            # text_file = create_dummy_text()
            # add_file_to_zip(zip_file, "Summary.txt", text_file)
            # pdf_file = create_dummy_pdf()
            # add_file_to_zip(zip_file, "Summary.pdf", pdf_file)
            # text_file = create_dummy_doc()
            # add_file_to_zip(zip_file, "Summary.docx", text_file)
        zip_buffer.seek(0)
        return StreamingResponse(
            zip_buffer,
            media_type="application/x-zip-compressed",
            headers={"Content-Disposition": "attachment; filename=files.zip"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e)) 

def add_folder_files_to_zip(zip_file, folder_path: str):
    """Simple function to read all files from a folder and add them to zip"""
    if os.path.exists(folder_path):
        for filename in os.listdir(folder_path):
            file_path = os.path.join(folder_path, filename)
            if os.path.isfile(file_path):
                with open(file_path, 'rb') as file:
                    zip_file.writestr(filename, file.read())

@router.post("/api/v1/upload")
async def upload_pdf():
# async def upload_pdf(pdf_file: UploadFile = File(...)):
    local_folder1: str = "Analysis"
    local_folder2: str = "Summary"
    local_folder3: str = "excel_data"
    score_json_path = "score.json"
    process_excel_files()
    try:
        import time
        # time.sleep(1)
        zip_buffer = BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
            # Add files from first sub-folder
            if os.path.exists(local_folder1):
                for filename in os.listdir(local_folder1):
                    file_path = os.path.join(local_folder1, filename)
                    if os.path.isfile(file_path):
                        with open(file_path, 'rb') as f:
                            zip_file.writestr(f"Analysis/{filename}", f.read())
            # Add files from second sub-folder
            if os.path.exists(local_folder2):
                for filename in os.listdir(local_folder2):
                    file_path = os.path.join(local_folder2, filename)
                    if os.path.isfile(file_path):
                        with open(file_path, 'rb') as f:
                            zip_file.writestr(f"Summary/{filename}", f.read())
            if os.path.exists(local_folder3):
                for filename in os.listdir(local_folder3):
                    file_path = os.path.join(local_folder3, filename)
                    if os.path.isfile(file_path):
                        with open(file_path, 'rb') as f:
                            zip_file.writestr(f"excel_data/{filename}", f.read())        
            # Add score.json file to zip
            if os.path.exists(score_json_path):
                with open(score_json_path, 'rb') as f:
                    zip_file.writestr("score.json", f.read())
        zip_buffer.seek(0)
        return StreamingResponse(
            zip_buffer,
            media_type="application/x-zip-compressed",
            headers={"Content-Disposition": "attachment; filename=files.zip"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))